"""Sequential Excel autopilot using the existing Studio engine and checkpoints."""
from pathlib import Path
import copy
import time
import uuid
import re
from .storage import atomic, read, redact
from .studio import prepare, project_name

def workbook(path, sheet=None):
    from openpyxl import load_workbook
    path = Path(path).expanduser().resolve()
    if not path.is_file() or path.suffix.lower() not in ('.xlsx', '.xlsm'):
        raise ValueError('Choose an existing .xlsx or .xlsm Excel file')
    if path.stat().st_size > 25_000_000:
        raise ValueError('Excel file must be smaller than 25 MB')
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        selected = sheet or book.sheetnames[0]
        if selected not in book.sheetnames:
            raise ValueError('Choose an existing worksheet')
        rows = []
        columns=None
        normalize=lambda value:re.sub(r'[^a-z0-9]+',' ',str(value or '').lower()).strip()
        title_headers={'title','story title','video title','project title'}
        script_headers={'script','story script','video script','one minute script','narration','story','full script'}
        for number, cells in enumerate(book[selected].iter_rows(max_col=50), 1):
            if number > 10000:
                raise ValueError('A worksheet may contain at most 10,000 rows')
            if not any(cell.value is not None and str(cell.value).strip() for cell in cells):continue
            if columns is None:
                labels=[normalize(cell.value) for cell in cells]
                titles=[i for i,label in enumerate(labels) if label in title_headers]
                scripts=[i for i,label in enumerate(labels) if label in script_headers]
                if titles or scripts:
                    if len(titles)!=1 or len(scripts)!=1:
                        raise ValueError(f'Row {number}: expected one Title / Story Title header and one Script / One-Minute Script header. Check the Excel column layout.')
                    columns=(titles[0],scripts[0]);continue
                columns=(0,1)
            chosen=[cells[i] for i in columns]
            title, script = [str(cell.value or '').strip() for cell in chosen]
            if not title and not script:
                continue
            if any(cell.data_type == 'f' for cell in chosen):
                raise ValueError(f'Row {number}: use plain text, not formulas, in title and script columns')
            if not title or not script:
                raise ValueError(f'Row {number}: both title and script are required')
            if re.fullmatch(r'\d+[.)]?',title) and columns==(0,1):
                raise ValueError(f'Row {number}: column A looks like a row number. Add Title and Script headers over the actual title/script columns, or use title in A and script in B.')
            rows.append({'row': number, 'title': title, 'script': script})
        if not rows:
            raise ValueError('This worksheet has no titles and scripts in columns A and B')
        from openpyxl.utils import get_column_letter
        return {'path': str(path), 'sheet': selected, 'sheets': book.sheetnames, 'rows': rows,
                'columns':{'title':get_column_letter(columns[0]+1),'script':get_column_letter(columns[1]+1)}}
    finally:
        book.close()

def prepare_batch(payload, settings, directory, occupied=(), animated=False):
    imported = workbook(payload.get('excel_path', ''), payload.get('excel_sheet'))
    options = copy.deepcopy(payload)
    for key in ('project', 'source_project', 'animate_source', 'scene_seed', 'scene_edit', 'package_seed', 'output', 'action', 'title', 'script', 'rows', 'checkpoint'):
        options.pop(key, None)
    options.update(auto_project=True, flow_url='', agent_tool_dir=options.get('agent_tool_dir') or str(Path(__file__).resolve().parents[1]/'animate_flow_backend'))
    if animated:
        options.update(provider='flow', browser_mode='cdp', share_google_profile=True, motion=False)
        # Generated references and their designs belong to one script/project.
        # User-selected reference_files remain selected for every workbook row.
        options.pop('references', None)
        options.pop('reference_descriptions', None)
    if options.get('provider') == 'flow':
        options.update(chrome_mode=True, share_google_profile=True)
    reserved = list(occupied)
    rows = []
    for row in imported['rows']:
        normalized = prepare({**options, **row, 'action': 'build'}, settings, reserved)
        reserved.append(normalized['project'])
        rows.append({**row, 'project': normalized['project'], 'state': 'pending', 'stage': 'plan'})
    return {**options, 'title': f'Excel batch · {imported["sheet"]} ({len(rows)} projects)', 'excel_path': imported['path'], 'excel_sheet': imported['sheet'], 'rows': rows,
            'checkpoint': str(Path(directory)/('animate_batches' if animated else 'batches')/(uuid.uuid4().hex+'.json'))}

def run(payload, settings, cancel, emit, execute, now=time.time):
    checkpoint = Path(payload['checkpoint'])
    saved = read(checkpoint)
    rows = saved.get('rows') or copy.deepcopy(payload['rows'])
    deadline = saved.get('wait_until', 0)
    def persist():
        atomic(checkpoint, {'rows': rows, 'wait_until': deadline})
    def notify(index, message, **extra):
        emit(message=message, progress=sum(r['state']=='completed' for r in rows)*100/len(rows),
             batch={'index': index, 'rows': rows, 'wait_until': deadline, **extra})
    def check():
        if cancel.is_set():
            raise RuntimeError('Batch stopped; Retry resumes saved work')
    failures = []
    for index, row in enumerate(rows):
        if row['state'] == 'completed':
            continue
        while now() < deadline:
            check()
            notify(index, f'Next project in {max(1,int(deadline-now()))} seconds')
            cancel.wait(min(1, max(0, deadline-now())))
        deadline = 0
        check()
        row.update(state='running', error='')
        persist()
        try:
            for stage in ('plan', 'build', 'export'):
                if ('plan','build','export').index(stage) < ('plan','build','export').index(row['stage']):
                    continue
                check()
                notify(index, f'Project {index+1}/{len(rows)} · {row["title"]} · {stage}')
                options = prepare({**payload, **row, 'action': stage,
                    'output': str(Path(row['project'])/(project_name(row['title'])+'.mp4'))}, settings)
                for key in ('rows','checkpoint','excel_path','excel_sheet','continue_on_error','state','stage','error','row','result'):
                    options.pop(key, None)
                def progress(**event):
                    event.pop('result', None)
                    event.pop('range_start', None)
                    event.pop('range_size', None)
                    if 'progress' in event:
                        event['progress'] = (index+min(99,max(0,float(event['progress'])))/100)*100/len(rows)
                    emit(**event)
                result = execute(options, copy.deepcopy(settings), cancel, progress)
                if result.get('review_required'):
                    raise RuntimeError('Cast review required. Open '+str(row['project'])+' in AI Studio, review and approve its cast, then retry this Excel batch. No images were submitted.')
                if stage == 'build' and (not result.get('video') or not Path(result['video']).is_file()):
                    raise RuntimeError('The project did not produce a rendered video; see the engine log')
                row['stage'] = {'plan':'build','build':'export','export':'export'}[stage]
                row['result'] = result
                persist()
            row['state'] = 'completed'
        except Exception as error:
            row.update(state='interrupted' if cancel.is_set() else 'failed', error=redact(str(error), settings))
            failures.append(row['title'])
            persist()
            notify(index, f'Project {index+1} failed: {error}')
            if cancel.is_set() or 'Account limit reached:' in str(error) or not payload.get('continue_on_error'):
                raise
        deadline = now()+60 if index < len(rows)-1 else 0
        persist()
        notify(index, f'Project {index+1}/{len(rows)} · {row["state"]}')
    emit(result={'checkpoint': str(checkpoint), 'completed': sum(r['state']=='completed' for r in rows), 'total': len(rows)})
    if failures:
        raise RuntimeError('Some projects failed. Retry resumes these rows: '+', '.join(failures))
